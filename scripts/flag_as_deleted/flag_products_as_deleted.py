# ============================================================
# TOURPLAN NX — OCULTAR PRODUCT CODES (Flag Product as Deleted)
# Copia adaptada para tp-nx-app (app local). El original para
# Google Colab vive sin cambios en el repo Flag-as-deleted.
# Mismos cambios que copy_products.py: config por variables de
# entorno, Chrome delegado a common/chrome_bootstrap.py, sin
# --headless/puerto fijo, perfil de Chrome aislado, log en
# carpeta temporal del OS, fix del import de importlib.util.
# El resto (selectores, login, logica de negocio) es igual.
#
# ADVERTENCIA: este script corre contra Tourplan de PRODUCCIÓN,
# no hay ambiente de test para este flujo. MODO default = lectura.
# ------------------------------------------------------------
#   VERSION : 1.1
#   FECHA   : 2026-07-23
#   v1.1: SUPPLIER_CODE ahora es filtro de búsqueda (autocomplete, como Location),
#         glosario fijo de SERVICE_TYPE, manejo del campo CLASS antes de Save, y
#         espera del botón SAVE por polling activo — según prompt actualizado con
#         grabaciones del 22/7.
# ============================================================

VERSION       = "1.1"
VERSION_FECHA = "2026-07-23"

# ── CONFIGURACIÓN — via variables de entorno (con default = valor original) ──
import os

MODO       = os.environ.get("TOURPLAN_MODO", "lectura")     # "lectura" | "escritura" | "completo"
EXCEL_PATH = os.environ.get("TOURPLAN_EXCEL_PATH", "/content/input_productos.xlsx")
HOJA       = os.environ.get("TOURPLAN_HOJA", "PRODUCTOS")

USERNAME   = os.environ.get("TOURPLAN_USERNAME", "poner minusculas")
PASSWORD   = os.environ.get("TOURPLAN_PASSWORD", "password")
BASE_URL   = os.environ.get("TOURPLAN_BASE_URL", "https://tourplannx.eurotur.com.ar/tourplannx")

SS_DIR     = os.environ.get("TOURPLAN_SS_DIR", "/content/screenshots")

FLAG_LABEL_TEXT = "flag product as deleted"

# ── PASO 0: Entorno ──────────────────────────────────────────
import sys, subprocess, importlib.util, shutil, time, traceback, re
from datetime import datetime

_t_inicio = time.time()

print("🔧 Verificando entorno...\n")

# 0.1 Paquetes Python
_PIPS_NEEDED = {
    "selenium":          "selenium",
    "openpyxl":          "openpyxl",
    "webdriver_manager": "webdriver-manager",
}
_faltantes = [pkg for mod, pkg in _PIPS_NEEDED.items()
              if importlib.util.find_spec(mod) is None]
if _faltantes:
    print(f"  ⏳ pip install {' '.join(_faltantes)} ...")
    subprocess.run([sys.executable, "-m", "pip", "install", "-q"] + _faltantes, check=True)
    print("  ✅ Paquetes OK")
else:
    print("  ✅ Paquetes ya instalados")

# 0.2 Google Chrome (multiplataforma — ver common/chrome_bootstrap.py)
from common.chrome_bootstrap import find_or_prepare_chrome
from common.abort import chequear_abort, AbortadoPorUsuario, ABORT_EXIT_CODE

CHROMIUM_BIN, ver_chrome = find_or_prepare_chrome()

# 0.3 Imports
from webdriver_manager.chrome import ChromeDriverManager
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException
import openpyxl

os.makedirs(SS_DIR, exist_ok=True)

# ── Constantes de espera ─────────────────────────────────────
WAIT_SHORT  = 5
WAIT_MEDIUM = 15
WAIT_LONG   = 30

# ── Selectores ───────────────────────────────────────────────
SEL = {
    "NAV_IMG"             : "nav img",
    "MENU_ITEMS"          : "div.nav-quick li label",
    "SEARCH_WRAPPER"      : "#searchWrapper",
    "SEARCH_BTN"          : "#searchWrapper li:nth-of-type(2) button",
    "FIELD_LOCATION"      : "div.parameters1 li:nth-of-type(1) input",
    "FIELD_SUPPLIER"      : "div.parameters1 li:nth-of-type(2) input",
    "SUGGEST_SELECTED_ROW": "tr.selectedRow > td.description",
    "SUGGEST_OPTION_CELL" : "td.description",
    "FIELD_OPTION"        : "#option tp-validator input",
    "FIELD_CLASS"         : "#tabs-product ul:nth-of-type(2) > li:nth-of-type(2) input",
    "TAB_RESULTS"         : "#tptablabel-productSearchResults",
    "RESULTS_CONTAINER"   : "#productSearchResults",
    "RESULTS_ROWS"        : "#productSearchResults tr",
    "RESULTS_SUPPLIER_COL": "td.tpcol-suppliercode",
    "TABS_PRODUCT"        : "#tabs-product",
    "SAVE_BTN"            : "tp-button.save > button",
}

# Glosario fijo Excel -> texto a buscar en el panel de Service Type (del prompt
# actualizado). Si el código de la celda no está acá, se usa el valor de la celda
# tal cual como texto de búsqueda (categorías fuera de estas 14 + 12, donde el
# propio Excel ya trae el texto listo).
#
# Las 12 siglas de abajo confirmadas por la usuaria (2026-08-18) aparecen en el
# panel con el prefijo "Z*NO USAR* - <SIGLA> - <DESCRIPCIÓN>" — "NO USAR" es una
# convención de organización interna de la empresa (agrupa estas categorías al
# fondo del listado), NO significa que estén deprecadas o que no deban tildarse.
# Se mapean a "<SIGLA> - <DESCRIPCIÓN>" (más específico que solo la sigla o solo
# la descripción) por dos motivos: evita depender de que sea la primera
# coincidencia dentro del ítem, y evita colisión con categorías reales que
# comparten una palabra — ej. TF="Transfer" y TN="Transfer Non-Accom" comparten
# la palabra "Transfer"; EX="Excursion" y EN="Excursion Non-Accom" comparten
# "Excursion". select_service_type() además excluye los ítems "Z*NO USAR*" al
# buscar una sigla que NO es una de estas 12, precisamente para no caer en esa
# colisión (ver comentario ahí).
Z_NO_USAR_SIGLAS = {"GA", "TK", "TN", "CH", "CM", "CO", "EN", "GU", "PJ", "ST", "TR", "TI"}

SERVICE_TYPE_GLOSARIO = {
    "HT": "Hotels",
    "HX": "Hotel Extras",
    "TF": "Transfer",
    "EX": "Excursion",
    "ML": "Meals",
    "RT": "Rental Car",
    "CR": "Cruceros",
    "FT": "Flight Ticket",
    "OC": "Operational Costs",
    "LN": "Location Name",
    "LP": "Location Pais",
    "MS": "Miscelaneos",
    "TA": "Travel Assistance",
    "PR": "Percepciones",
    # Grupo "Z*NO USAR*" (organización interna, ver comentario arriba) —
    # confirmadas por la usuaria el 2026-08-18.
    "GA": "GA - GASTRONOMIA",
    "TK": "TK - TICKETS",
    "TN": "TN - TRANSFER NON-ACCOM",
    "CH": "CH - CHOFER FREE LANCE PCM",
    "CM": "CM - CONSUMICIÓN",
    "CO": "CO - COSTO PCM",
    "EN": "EN - EXCURSION NON-ACCOM",
    "GU": "GU - GUIDE SERVICES",
    "PJ": "PJ - PEAJES",
    "ST": "ST - SERVICIOS DE TERCEROS",
    "TR": "TR - TRANSPORT",
    "TI": "TI - TIPS",
}

# ── Columnas Excel ────────────────────────────────────────────
COL_LOCATION      = "LOCATION"
COL_SERVICE_TYPE  = "SERVICE_TYPE"
COL_SUPPLIER_CODE = "SUPPLIER_CODE"
COL_PRODUCT_CODE  = "PRODUCT_CODE"
COL_ESTADO        = "ESTADO"
COL_OBSERVACIONES = "OBSERVACIONES"

# ── Helpers genéricos ─────────────────────────────────────────
_ss_n = [0]

def ss(driver, nombre):
    """`nombre` puede incluir el label de fila (ej. 'BUE/GU/1ALD03/PRUEBA'),
    con '/' que si se deja tal cual se interpreta como separador de carpeta en
    la ruta del archivo -> save_screenshot intenta escribir en subcarpetas
    inexistentes, falla y devuelve False SIN lanzar excepción, y el código
    anterior no chequeaba ese resultado (imprimía '📸 ...' igual, con solo el
    tramo final del path como nombre). Se sanea a [A-Za-z0-9_-] y se valida
    el resultado de save_screenshot."""
    _ss_n[0] += 1
    nombre_seguro = re.sub(r"[^A-Za-z0-9_-]+", "_", nombre)[:60]
    p = f"{SS_DIR}/{_ss_n[0]:03d}_{nombre_seguro}_{int(time.time())}.png"
    try:
        ok = driver.save_screenshot(p)
        if not ok:
            print(f"  ⚠️ No se pudo guardar captura: {os.path.basename(p)}")
            return None
        print(f"  📸 {os.path.basename(p)}")
        return p
    except Exception as e:
        print(f"  ⚠️ Error guardando captura ({e})")
        return None

def jc(driver, el):
    driver.execute_script("arguments[0].click();", el)

def set_val(driver, el, value, blur=False):
    driver.execute_script("""
        var inp = arguments[0], val = arguments[1], doBlur = arguments[2];
        inp.focus();
        var setter = Object.getOwnPropertyDescriptor(
            window.HTMLInputElement.prototype, 'value').set;
        setter.call(inp, val);
        inp.dispatchEvent(new Event('input',  {bubbles:true}));
        inp.dispatchEvent(new Event('change', {bubbles:true}));
        if (doBlur) {
            inp.blur();
            inp.dispatchEvent(new Event('blur', {bubbles:true}));
        }
    """, el, value, blur)

def wait(driver, css, t=WAIT_MEDIUM):
    return WebDriverWait(driver, t).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, css)))

def wait_click(driver, css, t=WAIT_MEDIUM):
    return WebDriverWait(driver, t).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, css)))

# ── Driver ────────────────────────────────────────────────────
def crear_driver():
    import tempfile

    opts = Options()
    # Sin --headless: esto corre en la PC de la persona (con pantalla), no en
    # el contenedor sin pantalla de Colab. Ademas, varias empresas bloquean
    # el modo headless de Chrome por politica de seguridad.
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_argument("--window-size=1704,1012")
    opts.add_argument("--disable-gpu")
    # Sin --remote-debugging-port fijo: chromedriver ya maneja su propio
    # puerto; un puerto fijo puede chocar con otro Chrome ya abierto.

    # Perfil temporal y aislado para esta corrida — evita que Chrome
    # "rebote" hacia una ventana ya abierta con los perfiles reales de la
    # persona (instancia unica de Windows compartiendo el mismo
    # user-data-dir), que Selenium no controla.
    _profile_dir = tempfile.mkdtemp(prefix="tourplan_chrome_profile_")
    opts.add_argument(f"--user-data-dir={_profile_dir}")
    opts.add_argument("--no-first-run")
    opts.add_argument("--no-default-browser-check")

    if CHROMIUM_BIN:
        opts.binary_location = CHROMIUM_BIN
        print(f"Chrome binary: {CHROMIUM_BIN}  ({ver_chrome})")
    else:
        raise RuntimeError("No se encontró Chrome funcional.")

    log_path = os.path.join(tempfile.gettempdir(), "chromedriver.log")
    try:
        drv_path = ChromeDriverManager().install()
        print(f"Chromedriver: {drv_path}")
        svc = Service(executable_path=drv_path, log_output=log_path)
        d = webdriver.Chrome(service=svc, options=opts)
        print("✅ Driver iniciado")
        return d
    except Exception as e:
        if os.path.exists(log_path):
            with open(log_path) as f:
                print(f"\n--- ChromeDriver log ---\n{f.read()[-3000:]}\n---")
        raise RuntimeError(f"No se pudo iniciar Chrome.\nError: {e}")

# ── Login ─────────────────────────────────────────────────────
def login(driver):
    print("🔐 Login...")
    driver.get(f"{BASE_URL}/#/login")
    time.sleep(6)
    ss(driver, "login_page")

    def _campos_visibles():
        return driver.execute_script("""
            function vis(e){return !!(e && (e.offsetWidth||e.offsetHeight
                                      ||e.getClientRects().length)
                                      && !e.disabled);}
            var txt = Array.from(document.querySelectorAll(
                "input[type='text'], input:not([type])")).filter(vis);
            var pwd = Array.from(document.querySelectorAll(
                "input[type='password']")).filter(vis);
            return [txt[0]||null, pwd[0]||null];
        """)

    u_el = p_el = None
    for _ in range(15):
        u_el, p_el = _campos_visibles()
        if u_el and p_el:
            break
        time.sleep(2)
    if not (u_el and p_el):
        ss(driver, "login_sin_campos")
        raise Exception("No aparecieron los campos de login")

    def _set(el, valor):
        try: el.clear()
        except Exception: pass
        try: el.click()
        except Exception: pass
        try:
            el.send_keys(valor)
        except Exception:
            driver.execute_script("""
                var el=arguments[0], v=arguments[1];
                var s=Object.getOwnPropertyDescriptor(
                    window.HTMLInputElement.prototype,'value').set;
                s.call(el,v);
                el.dispatchEvent(new Event('input',{bubbles:true}));
                el.dispatchEvent(new Event('change',{bubbles:true}));
            """, el, valor)

    _set(u_el, USERNAME)
    _set(p_el, PASSWORD)
    time.sleep(0.5)

    clic = driver.execute_script("""
        function vis(e){return !!(e && (e.offsetWidth||e.offsetHeight
                                  ||e.getClientRects().length) && !e.disabled);}
        var b = Array.from(document.querySelectorAll(
            "button.login, button[type='submit'], button")).filter(vis)
            .find(function(x){return /log\\s*in|ingresar|entrar|sign\\s*in/i
                                     .test((x.innerText||'')) ||
                                     x.classList.contains('login');});
        if (b){ b.click(); return (b.innerText||'button').trim(); }
        return null;
    """)
    if not clic:
        try: p_el.send_keys(Keys.ENTER)
        except Exception: pass

    time.sleep(8)
    assert "login" not in driver.current_url.lower(), "Login falló — verificá usuario/password"
    ss(driver, "post_login")
    print("✅ Login OK")

def _diagnostico(driver, etiqueta):
    try:
        url = driver.current_url
        titulo = driver.title
    except Exception:
        url = titulo = "?"
    print(f"  🩺 [{etiqueta}] url={url}  title={titulo}")
    ss(driver, f"diag_{etiqueta}")

# ── Navegación a Product Setup ────────────────────────────────
def open_product_setup(driver):
    print("Abriendo Product Setup...")
    driver.get(f"{BASE_URL}/#/product")
    time.sleep(4)
    try:
        wait(driver, SEL["SEARCH_WRAPPER"], t=WAIT_MEDIUM)
        print("  ✅ Navegación directa por hash OK")
        return
    except TimeoutException:
        _diagnostico(driver, "hash_directo_fallo")
        print("  ⚠ Hash directo no cargó, probando flujo de menú (Home → nav → Product Setup)...")

    driver.get(f"{BASE_URL}/#/home")
    time.sleep(3)
    try:
        nav_img = wait_click(driver, SEL["NAV_IMG"])
    except TimeoutException:
        _diagnostico(driver, "nav_img_no_encontrado")
        raise Exception("No se encontró el ícono de navegación ('nav img') en Home")
    jc(driver, nav_img)
    time.sleep(1)

    item = driver.execute_script("""
        var items = Array.from(document.querySelectorAll(arguments[0]));
        var target = items.find(function(l){
            return (l.innerText||l.textContent||'').trim().toLowerCase()
                   .indexOf('product setup') !== -1;
        });
        return target || null;
    """, SEL["MENU_ITEMS"])
    if not item:
        _diagnostico(driver, "menu_sin_product_setup")
        raise Exception("No se encontró 'Product Setup' en el menú de navegación")
    jc(driver, item)
    time.sleep(3)
    try:
        wait(driver, SEL["SEARCH_WRAPPER"], t=WAIT_LONG)
    except TimeoutException:
        _diagnostico(driver, "product_setup_no_cargo_via_menu")
        raise
    print("  ✅ Product Setup abierto vía menú")

def reset_search(driver):
    """Re-navegación completa Home → Product, igual que search_options() en
    copy_products.py — probado que limpia el estado anterior de forma
    confiable (filtros, service type y resultados previos)."""
    driver.get(f"{BASE_URL}/#/home")
    time.sleep(2)
    driver.get(f"{BASE_URL}/#/product")
    time.sleep(5)
    try:
        wait(driver, SEL["SEARCH_WRAPPER"], t=WAIT_MEDIUM)
    except TimeoutException:
        open_product_setup(driver)

def open_search_filters(driver):
    """El panel con los campos de búsqueda (Location/Service Type/Option) no está
    en el DOM hasta hacer click en el botón de búsqueda dentro de #searchWrapper."""
    try:
        btn = wait_click(driver, SEL["SEARCH_BTN"], t=WAIT_SHORT)
        jc(driver, btn)
        time.sleep(1)
    except TimeoutException:
        pass  # puede que el panel ya esté abierto (p.ej. tras un Save previo)

# ── Location / Supplier (autocomplete) ─────────────────────────
def _select_autocomplete(driver, field_css, value):
    """Patrón autocomplete compartido por Location y Supplier (Supplier confirmado
    como filtro de búsqueda en la grabación del 22/7 — antes solo se usaba para
    verificar la fila en los resultados): set_val + esperar tr.selectedRow >
    td.description y clickearla; si no aparece sugerencia, Enter como fallback
    (mismo patrón probado en copy_products.py)."""
    campo = wait(driver, field_css)
    campo.click()
    time.sleep(0.3)
    set_val(driver, campo, value)
    time.sleep(1.2)
    try:
        row = WebDriverWait(driver, WAIT_SHORT).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, SEL["SUGGEST_SELECTED_ROW"])))
        jc(driver, row)
        time.sleep(0.8)
        return True
    except TimeoutException:
        pass
    try:
        campo.send_keys(Keys.RETURN)
    except Exception:
        pass
    time.sleep(0.8)
    return True

def select_location(driver, location_code):
    return _select_autocomplete(driver, SEL["FIELD_LOCATION"], location_code)

def select_supplier(driver, supplier_code):
    """Selector confirmado en la grabación del 22/7: div.parameters1
    li:nth-of-type(2) input — mismo comportamiento autocomplete que Location."""
    return _select_autocomplete(driver, SEL["FIELD_SUPPLIER"], supplier_code)

# ── Service Type ──────────────────────────────────────────────
def select_service_type(driver, service_type_code):
    """Resuelve el código de la celda contra SERVICE_TYPE_GLOSARIO (ej. 'EX' ->
    'Excursion'); si no está en el glosario, usa el valor de la celda tal cual
    (categorías fuera de las 26 mapeadas, donde el Excel ya trae el texto a
    buscar). Busca el <li> del panel de filtros (#productSearchFilter) cuyo
    texto CONTENGA ese texto como PALABRA COMPLETA (delimitada por inicio/fin
    de texto o cualquier caracter no alfanumérico), sin distinguir mayúsculas/
    minúsculas.

    Exige palabra completa a propósito: un substring suelto sobre una sigla
    sin glosario puede matchear por accidente dentro de la palabra de otra
    categoría — confirmado en corrida real, 'TR' (antes de mapearse) tildaba
    "Hotel Extras" porque "Ex-TR-as" contiene 'tr'.

    Además, si `code` NO es una de las siglas "Z*NO USAR*" (Z_NO_USAR_SIGLAS),
    se descartan como candidatos los <li> cuyo texto contenga "no usar" —
    varias de esas categorías repiten una palabra de una categoría real (ej.
    TN="Transfer Non-Accom" comparte "Transfer" con TF; EN="Excursion
    Non-Accom" comparte "Excursion" con EX), y sin este descarte la búsqueda
    de la categoría real podía terminar seleccionando la de "Z*NO USAR*" por
    palabra compartida.

    Si la sigla no aparece como palabra completa en ningún ítem candidato del
    panel, se reporta 'no encontrado' y el llamador (buscar_productos) sigue
    con "All Services" en vez de tildar algo equivocado en silencio.
    Reintenta unas veces porque el panel puede tardar en poblarse."""
    code = service_type_code.strip()
    if not code:
        return True, None

    texto_buscar = SERVICE_TYPE_GLOSARIO.get(code.upper(), code)
    excluir_no_usar = code.upper() not in Z_NO_USAR_SIGLAS
    resultado = None
    for _ in range(3):
        resultado = driver.execute_script("""
            var needle = arguments[0].toLowerCase();
            var excluirNoUsar = arguments[1];
            function esPalabra(c) { return !!c && /[a-z0-9]/i.test(c); }
            var textos = [], elegido = null;
            for (var li of document.querySelectorAll('#productSearchFilter li')) {
                if (!li.offsetParent) continue;
                var t = (li.textContent || '').trim();
                if (!t || t.length > 80) continue;
                textos.push(t);
                if (elegido) continue;
                if (excluirNoUsar && t.toLowerCase().indexOf('no usar') !== -1) continue;
                var tl = t.toLowerCase();
                var idx = tl.indexOf(needle);
                while (idx !== -1) {
                    var antes = idx > 0 ? tl[idx - 1] : '';
                    var despues = (idx + needle.length < tl.length) ? tl[idx + needle.length] : '';
                    if (!esPalabra(antes) && !esPalabra(despues)) { elegido = li; break; }
                    idx = tl.indexOf(needle, idx + 1);
                }
            }
            if (elegido) {
                elegido.click();
                return {ok: true, texto: (elegido.textContent || '').trim().slice(0, 60)};
            }
            return {ok: false, textos: textos.slice(0, 40)};
        """, texto_buscar, excluir_no_usar)
        if resultado and resultado.get("ok"):
            break
        time.sleep(1.5)

    if resultado and resultado.get("ok"):
        time.sleep(1.2)
        print(f"  Service Type: {resultado.get('texto')}")
        return True, None

    textos = (resultado or {}).get("textos", [])
    print(f"  🩺 Service Type '{service_type_code}' (buscado como '{texto_buscar}'): no "
          f"encontrado. Ítems <li> visibles en #productSearchFilter (máx 40): {textos}")
    return False, (f"Service type '{service_type_code}' (buscado como '{texto_buscar}') "
                    f"no encontrado ({len(textos)} ítems <li> visibles en #productSearchFilter)")

# ── Product Code ──────────────────────────────────────────────
def enter_product_code(driver, product_code):
    """Selector e interacción tomados de copy_products.py (probado):
    '#option tp-validator input', set_val + Enter dispara la búsqueda —
    no existe un botón SEARCH separado en este panel."""
    code_input = wait(driver, SEL["FIELD_OPTION"])
    code_input.click()
    time.sleep(0.3)
    set_val(driver, code_input, product_code)
    time.sleep(0.5)
    try:
        code_input.send_keys(Keys.RETURN)
    except Exception:
        pass
    time.sleep(1)

# ── Resultados ────────────────────────────────────────────────
def get_result_rows(driver):
    try:
        tab = wait_click(driver, SEL["TAB_RESULTS"])
    except TimeoutException:
        _diagnostico(driver, "tab_results_no_encontrado")
        return []
    jc(driver, tab)
    time.sleep(2)
    try:
        wait(driver, SEL["RESULTS_CONTAINER"], t=WAIT_MEDIUM)
    except TimeoutException:
        _diagnostico(driver, "results_container_no_cargo")
        return []

    rows = driver.find_elements(By.CSS_SELECTOR, SEL["RESULTS_ROWS"])
    data_rows = [r for r in rows if r.find_elements(By.TAG_NAME, "td")]
    print(f"  🩺 Resultados: {len(rows)} <tr> en el contenedor, {len(data_rows)} con <td> (filas de datos)")
    if not data_rows:
        _diagnostico(driver, "results_sin_filas")
    return data_rows

def match_supplier_row(rows, supplier_code):
    """Columna confirmada por inspección real de la grilla: td.tpcol-suppliercode
    tiene el código exacto (ej. '1ALD03'), td.tpcol-supplierdescription tiene el
    nombre (ej. 'Aldaz Carlos (Guia)'). Se compara por igualdad exacta contra la
    columna de código. Si no hay exactamente un match, se vuelca clase+texto de
    cada celda de cada fila para poder diagnosticar (p.ej. código de proveedor
    incorrecto en el Excel, no un problema del selector)."""
    if len(rows) == 1:
        return rows[0], None

    target = str(supplier_code).strip().upper()
    matches = []
    dump = []
    for i, r in enumerate(rows):
        sup_cells = r.find_elements(By.CSS_SELECTOR, SEL["RESULTS_SUPPLIER_COL"])
        sup_text = sup_cells[0].text.strip() if sup_cells else ""
        cells = r.find_elements(By.TAG_NAME, "td")
        cell_info = [f"{c.get_attribute('class') or ''}={c.text.strip()!r}" for c in cells]
        dump.append(f"fila {i}: " + " | ".join(cell_info))
        if target and sup_text.upper() == target:
            matches.append(r)

    if len(matches) == 1:
        return matches[0], None

    print(f"  🩺 Supplier '{supplier_code}': {len(matches)} match(es) de {len(rows)} filas "
          f"por td.tpcol-suppliercode. Detalle de columnas:")
    for d in dump:
        print(f"      · {d}")
    muestra = " || ".join(dump)[:600]
    return None, (f"No se pudo identificar la fila correcta ({len(rows)} resultados, "
                   f"{len(matches)} matches por supplier). Columnas: {muestra}")

def _row_field(r, css_class):
    cells = r.find_elements(By.CSS_SELECTOR, f"td.{css_class}")
    return cells[0].text.strip() if cells else ""

def _row_key(r):
    """Clave estable para volver a encontrar la misma fila tras un re-búsqueda
    (usa el texto completo de la fila, no solo el proveedor, porque con más de
    un filtro vacío los resultados pueden variar también en location/producto)."""
    return "|".join(c.text.strip() for c in r.find_elements(By.TAG_NAME, "td"))

def _row_label(r):
    partes = [
        _row_field(r, "tpcol-locationcode"),
        _row_field(r, "tpcol-servicecode"),
        _row_field(r, "tpcol-suppliercode"),
        _row_field(r, "tpcol-optioncode"),
    ]
    return "/".join(p for p in partes if p) or "fila"

# ── Checkbox "Flag Product as Deleted" ────────────────────────
def _dump_checkbox_candidates(driver):
    """Cuando no se encuentra el checkbox por label, se vuelca tag+clase+texto
    de todo elemento candidato (label, input[type=checkbox], o cualquier tag/
    clase que contenga 'checkbox') dentro de #tabs-product, para poder
    identificar el selector real en el log (mismo enfoque que reveló
    td.tpcol-suppliercode antes)."""
    return driver.execute_script("""
        var container = document.querySelector(arguments[0]);
        if (!container) return [];
        var all = Array.from(container.querySelectorAll('*'));
        var out = [];
        for (var i = 0; i < all.length && out.length < 60; i++) {
            var el = all[i];
            var tag = el.tagName.toLowerCase();
            var cls = (el.className && el.className.toString) ? el.className.toString() : '';
            var esCandidato = tag === 'label' || tag === 'input' ||
                               tag.indexOf('checkbox') !== -1 ||
                               cls.toLowerCase().indexOf('checkbox') !== -1;
            if (!esCandidato) continue;
            var txt = (el.innerText || el.textContent || '').trim().replace(/\\s+/g, ' ').slice(0, 50);
            out.push(tag + '.' + cls + '=' + JSON.stringify(txt));
        }
        return out;
    """, SEL["TABS_PRODUCT"])

# Confirmado por volcado real de #tabs-product: el checkbox NO está envuelto
# por su label (como se asumía originalmente) — son hermanos independientes:
#   <tp-checkbox><label class="tpcheckbox-indent tpcheckbox">
#     <input class="tpcheckbox" type="checkbox"><span class="tpcheckboxtext">
#   </label></tp-checkbox>
#   <label class="default tplabel tplabel-checkbox tplabel-flagasdeleted">
#     FLAG PRODUCT AS DELETED</label>
# Se ubica el label descriptivo por su clase específica (más estable que el
# texto) y se busca el input.tpcheckbox subiendo por los ancestros comunes.
_FIND_FLAG_INPUT_JS = """
    var labelText = arguments[0];
    var container = document.querySelector(arguments[1]);
    if (!container) return {found:false, reason:'no #tabs-product'};

    var label = container.querySelector('label.tplabel-flagasdeleted') ||
        Array.from(container.querySelectorAll('label')).find(function(l){
            return (l.innerText||l.textContent||'').trim().toLowerCase()
                   .indexOf(labelText) !== -1;
        });
    if (!label) return {found:false, reason:"checkbox 'Flag Product as Deleted' no encontrado"};

    var input = null, scope = label;
    for (var depth = 0; depth < 4 && !input; depth++) {
        scope = scope.parentElement;
        if (!scope) break;
        input = scope.querySelector('input.tpcheckbox, input[type="checkbox"]');
    }
    if (!input) input = container.querySelector('input.tpcheckbox, input[type="checkbox"]');
    if (!input) return {found:false, reason:'label encontrado pero no el input del checkbox asociado'};
"""

def read_flag_checkbox(driver):
    result = driver.execute_script(_FIND_FLAG_INPUT_JS + """
        return {found:true, checked: !!input.checked};
    """, FLAG_LABEL_TEXT, SEL["TABS_PRODUCT"])
    if not result or not result.get("found"):
        dump = _dump_checkbox_candidates(driver)
        if dump:
            print(f"  🩺 Checkbox no encontrado por label. Candidatos en #tabs-product:")
            for d in dump:
                print(f"      · {d}")
            muestra = " || ".join(dump)[:600]
            reason = (result or {}).get("reason", "checkbox no encontrado")
            return {"found": False, "reason": f"{reason}. Candidatos: {muestra}"}
    return result

def click_flag_checkbox(driver):
    result = driver.execute_script(_FIND_FLAG_INPUT_JS + """
        input.click();
        return {found:true};
    """, FLAG_LABEL_TEXT, SEL["TABS_PRODUCT"])
    if not result or not result.get("found"):
        return False, (result or {}).get("reason", "checkbox no encontrado")
    time.sleep(0.6)
    return True, None

# ── Campo CLASS (bloquea el Save si queda vacío) ────────────────
def _click_first_suggestion(driver, timeout=WAIT_SHORT):
    """Clickea la primera fila visible de una tabla de sugerencias/opciones (mismo
    componente que usan los autocomplete de Location/Supplier), sin asumir un texto
    fijo — el prompt pide explícitamente no hardcodear 'Unassigned' como primera
    opción de CLASS, por si en otros casos el orden difiere."""
    try:
        WebDriverWait(driver, timeout).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, SEL["SUGGEST_OPTION_CELL"])))
    except TimeoutException:
        return False
    opciones = [c for c in driver.find_elements(By.CSS_SELECTOR, SEL["SUGGEST_OPTION_CELL"])
                if c.is_displayed()]
    if not opciones:
        return False
    jc(driver, opciones[0])
    return True

def ensure_class_field(driver):
    """Paso nuevo del prompt (grabación del 22/7): el campo CLASS
    (#tabs-product ul:nth-of-type(2) > li:nth-of-type(2) input) bloquea el Save si
    queda vacío. Si está vacío se abre su dropdown y se clickea la primera opción
    (sin asumir 'Unassigned'); si ya tiene un valor cargado, no se toca (es un dato
    de negocio existente). Si el campo no existe en este producto/tab, no bloquea."""
    try:
        class_input = wait(driver, SEL["FIELD_CLASS"], t=WAIT_SHORT)
    except TimeoutException:
        return None
    valor_actual = (class_input.get_attribute("value") or "").strip()
    if valor_actual:
        return None
    class_input.click()
    time.sleep(0.3)
    if not _click_first_suggestion(driver):
        return "CLASS estaba vacío y no aparecieron opciones para completarlo"
    time.sleep(0.8)
    return None

# ── Botón SAVE (espera activa de habilitado, no sleep fijo) ────
SAVE_POLL_INTERVAL = 0.4
SAVE_POLL_TIMEOUT  = 15  # seg — pedido explícito del prompt (10-15s), con polling ~300-500ms

def wait_save_habilitado(driver, timeout=SAVE_POLL_TIMEOUT, interval=SAVE_POLL_INTERVAL):
    """Polling activo del estado habilitado del SAVE — el prompt pide explícitamente no
    usar un sleep fijo (poco confiable) sino revisar cada ~300-500ms si se habilitó.
    Distingue 'no existe en el DOM' de 'existe pero nunca se habilitó' (p.ej. CLASS
    vacío u otro campo obligatorio sin completar bloqueando el Save)."""
    fin = time.time() + timeout
    encontrado_alguna_vez = False
    while time.time() < fin:
        elems = driver.find_elements(By.CSS_SELECTOR, SEL["SAVE_BTN"])
        if elems:
            encontrado_alguna_vez = True
            try:
                if elems[0].is_displayed() and elems[0].is_enabled():
                    return elems[0], None
            except Exception:
                pass  # elemento stale por un re-render; se reintenta en la próxima vuelta
        time.sleep(interval)
    if not encontrado_alguna_vez:
        return None, "el botón SAVE no se encontró en la página"
    return None, "el botón SAVE nunca se habilitó (timeout)"

def _dump_save_candidates(driver):
    """Cuando `tp-button.save > button` no aparece clickeable, se vuelca
    tag+clase+texto+disabled de cualquier botón/tp-button cuyo texto o clase
    contenga 'save', en toda la página (no solo #tabs-product, por si el
    botón vive fuera de ese contenedor), para diagnosticar sin acceso al DOM."""
    return driver.execute_script("""
        var all = Array.from(document.querySelectorAll('button, tp-button'));
        var out = [];
        for (var i = 0; i < all.length && out.length < 30; i++) {
            var el = all[i];
            var cls = (el.className && el.className.toString) ? el.className.toString() : '';
            var txt = (el.innerText || el.textContent || '').trim().replace(/\\s+/g, ' ');
            var esCandidato = cls.toLowerCase().indexOf('save') !== -1 ||
                               txt.toLowerCase().indexOf('save') !== -1;
            if (!esCandidato) continue;
            var vis = !!(el.offsetWidth || el.offsetHeight || el.getClientRects().length);
            out.push(el.tagName.toLowerCase() + '.' + cls + '=' + JSON.stringify(txt.slice(0, 40)) +
                      ' visible=' + vis + ' disabled=' + !!el.disabled);
        }
        return out;
    """)

# ── Excel ─────────────────────────────────────────────────────
def load_excel(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[HOJA]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {h: i + 1 for i, h in enumerate(headers) if h}
    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        d = dict(zip(headers, row))
        d["__row_idx__"] = row_idx
        rows.append(d)
    print(f"Excel cargado: {len(rows)} fila(s) en '{HOJA}'")
    return wb, rows, col_idx

def update_row(wb, row_idx, col_idx, path, estado=None, observaciones=None):
    ws = wb[HOJA]
    if estado is not None:
        ws.cell(row=row_idx, column=col_idx[COL_ESTADO]).value = estado
    if observaciones is not None:
        ws.cell(row=row_idx, column=col_idx[COL_OBSERVACIONES]).value = observaciones
    wb.save(path)

def _con_avisos(avisos, msg=None):
    partes = list(avisos)
    if msg:
        partes.append(msg)
    return " | ".join(partes) if partes else None

# ── Búsqueda completa (Location + Supplier + Service Type + Product Code) ──
def buscar_productos(driver, location, supplier_code, service_type, product_code):
    """Ejecuta una búsqueda completa desde cero (reset + abrir filtros + Location +
    Supplier + Service Type + Product Code, en ese orden) y devuelve (rows, aviso,
    error). SUPPLIER_CODE ahora también se usa como filtro de búsqueda cuando está
    completo (antes solo servía para verificar la fila en los resultados). Si algún
    campo viene vacío, no se completa ese filtro y la búsqueda toma todos los
    resultados para los filtros que sí están completos (igual que search_options()
    en copy_products.py). `rows` es [] si la búsqueda no encontró nada; `error` solo
    se usa para fallas estructurales (el panel o el campo de Location no cargaron)."""
    reset_search(driver)
    open_search_filters(driver)

    try:
        wait(driver, SEL["FIELD_LOCATION"], t=WAIT_MEDIUM)
    except TimeoutException:
        return None, None, "El panel de búsqueda (Location/Service Type/Option) no cargó"

    if location:
        if not select_location(driver, location):
            return None, None, f"No se pudo seleccionar location '{location}' en la búsqueda"

    if supplier_code:
        select_supplier(driver, supplier_code)

    aviso = None
    if service_type:
        ok, msg = select_service_type(driver, service_type)
        if not ok:
            # No se tilda nada: queda "All Services" y se sigue con los demás filtros.
            aviso = f"Service type '{service_type}' no encontrado, se dejó 'All Services' ({msg})"
            print(f"  ⚠ {aviso}")

    if product_code:
        enter_product_code(driver, product_code)

    return get_result_rows(driver), aviso, None

def _click_target_cell(target_row):
    """`on-click-editaction` es una CLASE CSS (no un atributo con valor —
    el dump anterior solo imprimía class+texto, lo que hacía parecer un
    atributo='valor') que marca qué celdas abren el detalle al clickear.
    La primera celda de la fila viene vacía (class='', probablemente un
    espaciador) y NO tiene esa clase, por eso clickearla no hacía nada."""
    cells = target_row.find_elements(By.TAG_NAME, "td")
    for c in cells:
        clases = (c.get_attribute("class") or "").split()
        if "on-click-editaction" in clases:
            return c
    return cells[0] if cells else None

# ── Abrir un producto de la grilla y tildar/leer el checkbox ───────────────
def procesar_item(driver, target_row, location, service_type, product_code, label):
    """Clickea la fila ya localizada en la grilla de resultados y aplica el
    flujo de lectura o de tildado+guardado sobre ESE producto puntual.
    Devuelve (estado_item, detalle_item) — no es el ESTADO final de la fila
    del Excel, que puede agrupar varios ítems si SUPPLIER_CODE está vacío."""
    click_target = _click_target_cell(target_row) or target_row
    jc(driver, click_target)
    time.sleep(1.5)

    try:
        wait(driver, SEL["TABS_PRODUCT"], t=WAIT_LONG)
    except TimeoutException:
        _diagnostico(driver, f"detalle_no_cargo_{label or 'item'}")
        return "ERROR", "el detalle del producto no cargó (#tabs-product)"

    etiqueta_archivo = f"{location}_{service_type}_{product_code}_{label or 'item'}"

    if MODO == "lectura":
        info = read_flag_checkbox(driver)
        ss(driver, f"lectura_{etiqueta_archivo}")
        if not info or not info.get("found"):
            return "PENDIENTE", f"[LECTURA] {info.get('reason') if info else 'checkbox no encontrado'}"
        return "PENDIENTE", (f"[LECTURA] checkbox actualmente "
                              f"{'tildado' if info.get('checked') else 'sin tildar'}")

    ok, msg = click_flag_checkbox(driver)
    if not ok:
        return "ERROR", f"no se encontró o no se pudo tildar 'Flag Product as Deleted' ({msg})"

    info = read_flag_checkbox(driver)
    if not info or not info.get("found") or not info.get("checked"):
        return "ERROR", "no quedó tildado 'Flag Product as Deleted'"

    class_err = ensure_class_field(driver)
    if class_err:
        ss(driver, f"class_vacio_{label or 'item'}")
        return "ERROR", class_err

    ss(driver, etiqueta_archivo)

    save_btn, err = wait_save_habilitado(driver)
    if err:
        _diagnostico(driver, f"save_no_habilitado_{label or 'item'}")
        dump = _dump_save_candidates(driver)
        detalle = err
        if dump:
            print(f"  🩺 Botón SAVE no habilitado. Candidatos en la página:")
            for d in dump:
                print(f"      · {d}")
            detalle += f". Candidatos: {' || '.join(dump)[:600]}"
        return "ERROR", detalle
    jc(driver, save_btn)

    time.sleep(3)
    return "HECHO", None

# ── Procesamiento de una fila ──────────────────────────────────
def process_row(driver, row):
    location      = str(row.get(COL_LOCATION) or "").strip()
    service_type  = str(row.get(COL_SERVICE_TYPE) or "").strip()
    supplier_code = str(row.get(COL_SUPPLIER_CODE) or "").strip()
    product_code  = str(row.get(COL_PRODUCT_CODE) or "").strip()
    row_idx       = row["__row_idx__"]

    if not (location or service_type or product_code):
        return "ERROR", ("Debe completarse al menos uno de LOCATION, SERVICE_TYPE o "
                          "PRODUCT_CODE (no pueden estar los tres vacíos — se procesaría "
                          "todo el catálogo sin filtro alguno)")

    print(f"  🔍 Location={location or '(vacío → todas)'}  "
          f"ServiceType={service_type or '(vacío → todos)'}  "
          f"Supplier={supplier_code or '(vacío → todos)'}  "
          f"Product={product_code or '(vacío → todos)'}")

    avisos = []

    rows, aviso, err = buscar_productos(driver, location, supplier_code, service_type, product_code)
    if err:
        _diagnostico(driver, f"busqueda_fallo_row{row_idx}")
        return "ERROR", err
    if aviso:
        avisos.append(aviso)

    if not rows:
        return "ERROR", _con_avisos(avisos, "Producto no encontrado")

    # ── SUPPLIER_CODE completo: identificar y procesar una sola fila ──
    if supplier_code:
        matched_row, msg = match_supplier_row(rows, supplier_code)
        if matched_row is None:
            return "ERROR", _con_avisos(avisos, msg)
        estado_item, detalle_item = procesar_item(
            driver, matched_row, location, service_type, product_code, supplier_code)
        return estado_item, _con_avisos(avisos, detalle_item)

    # ── SUPPLIER_CODE (u otro filtro) vacío: procesar TODOS los resultados ──
    # Se identifica cada fila por su contenido completo (no solo el proveedor),
    # porque con más de un filtro vacío los resultados también pueden variar
    # en location o en product code, no solo en supplier.
    filas_unicas = []
    claves_vistas = set()
    for r in rows:
        clave = _row_key(r)
        if clave not in claves_vistas:
            claves_vistas.add(clave)
            filas_unicas.append((clave, _row_label(r)))

    print(f"  ℹ️ Filtro(s) vacío(s): se van a procesar los {len(filas_unicas)} "
          f"resultado(s) encontrados: {[lbl for _, lbl in filas_unicas]}")

    resumen = []
    n_error = 0
    for i, (clave, label) in enumerate(filas_unicas):
        if i == 0:
            rows_actuales = rows
        else:
            rows_actuales, _aviso_ignorado, err2 = buscar_productos(
                driver, location, supplier_code, service_type, product_code)
            if err2:
                resumen.append(f"{label}: ERROR re-búsqueda falló ({err2})")
                n_error += 1
                continue
            if not rows_actuales:
                resumen.append(f"{label}: no apareció en la re-búsqueda "
                                f"(¿ya estaba tildado y desapareció del listado?)")
                continue

        objetivo = next((r for r in rows_actuales if _row_key(r) == clave), None)
        if objetivo is None:
            resumen.append(f"{label}: no encontrado en la re-búsqueda")
            continue

        estado_item, detalle_item = procesar_item(
            driver, objetivo, location, service_type, product_code, label)
        resumen.append(f"{label}: {estado_item}"
                        + (f" ({detalle_item})" if detalle_item else ""))
        if estado_item == "ERROR":
            n_error += 1

    estado_final = "ERROR" if n_error else ("PENDIENTE" if MODO == "lectura" else "HECHO")
    return estado_final, _con_avisos(avisos, " || ".join(resumen))

# ── MAIN ──────────────────────────────────────────────────────
print("=" * 60)
print(f"  OCULTAR PRODUCT CODES  v{VERSION}  [{VERSION_FECHA}]  MODO={MODO}")
print("=" * 60)

if MODO not in ("lectura", "escritura", "completo"):
    raise ValueError(f"MODO inválido: '{MODO}'. Usar 'lectura', 'escritura' o 'completo'.")

if MODO == "lectura":
    print("ℹ️  MODO LECTURA: no se tildará ni guardará nada. Solo se valida búsqueda,\n"
          "   matching de fila y estado actual del checkbox.")

if not os.path.exists(EXCEL_PATH):
    raise FileNotFoundError(f"No encontré el Excel en {EXCEL_PATH} — subilo a /content/")

wb, rows, col_idx = load_excel(EXCEL_PATH)
pendientes = [r for r in rows
              if str(r.get(COL_ESTADO) or "").strip().upper() == "PENDIENTE"]

print(f"Filas PENDIENTE: {len(pendientes)}")
for f in pendientes:
    print(f"  · fila {f['__row_idx__']}: {f.get(COL_LOCATION)} / {f.get(COL_SERVICE_TYPE)} / "
          f"{f.get(COL_SUPPLIER_CODE)} / {f.get(COL_PRODUCT_CODE)}")

if not pendientes:
    print("\n⛔ Sin filas PENDIENTE. Verificá que la columna ESTADO tenga 'PENDIENTE'.")
    raise SystemExit

driver = crear_driver()
_abortado = False
try:
    login(driver)
    open_product_setup(driver)

    for row in pendientes:
        chequear_abort()
        row_idx = row["__row_idx__"]
        print(f"\n{'─'*60}")
        print(f"Fila {row_idx}: {row.get(COL_LOCATION)} / {row.get(COL_SERVICE_TYPE)} / "
              f"{row.get(COL_SUPPLIER_CODE)} / {row.get(COL_PRODUCT_CODE)}")

        estado, observaciones = "ERROR", "Error desconocido"
        try:
            estado, observaciones = process_row(driver, row)
        except Exception:
            estado = "ERROR"
            observaciones = traceback.format_exc(limit=3)
            ss(driver, f"fatal_row{row_idx}")

        print(f"  Estado: {estado}" + (f" — {observaciones}" if observaciones else ""))
        update_row(wb, row_idx, col_idx, EXCEL_PATH, estado=estado, observaciones=observaciones)

except AbortadoPorUsuario:
    _abortado = True
    print("\n⏸️  Corrida abortada por el usuario — las filas que no llegaron a "
          "procesarse quedan en PENDIENTE para retomar en otra corrida.")

finally:
    driver.quit()
    dur = int(time.time() - _t_inicio)
    m, s = divmod(dur, 60)
    print(f"\n🏁 Fin. Duración: {m}m {s:02d}s")
    print(f"📄 Excel: {EXCEL_PATH}")
    print(f"📂 Screenshots: {SS_DIR}")

if _abortado:
    sys.exit(ABORT_EXIT_CODE)

try:
    from google.colab import files
    files.download(EXCEL_PATH)
    print("📥 Excel descargado.")
except Exception:
    pass
